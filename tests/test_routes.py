from pathlib import Path
from io import BytesIO
from unittest.mock import MagicMock, patch
import time

from docx import Document

from app import create_app
from app.folder_ingest import JOB_STORE


def test_analyze_accepts_uploaded_txt_file(tmp_path: Path):
    app = create_app({"TESTING": True, "OPINION_DATA_DIR": tmp_path / "missing"})

    with app.test_client() as client:
        response = client.post(
            "/analyze",
            data={
                "opinion_file": (
                    BytesIO(b"Opinion by Judge Jane A. Doe\n\nThe motion was granted by the court."),
                    "upload.txt",
                ),
            },
            content_type="multipart/form-data",
        )

    assert response.status_code == 200
    assert b"Jane A. Doe" in response.data


def test_analyze_extracts_text_from_uploaded_pdf(tmp_path: Path):
    app = create_app({"TESTING": True, "OPINION_DATA_DIR": tmp_path / "missing"})

    pdf_file = MagicMock()
    pdf_file.filename = "brief.pdf"
    pdf_file.stream.read.return_value = b"%PDF-1.4 fake"

    fake_page = MagicMock()
    fake_page.extract_text.return_value = (
        "Opinion by Judge Robert T. Hill\n\nThe motion was denied because the argument was unclear."
    )

    fake_reader = MagicMock()
    fake_reader.pages = [fake_page]

    with app.test_client() as client, patch("app.routes.PdfReader", return_value=fake_reader):
        response = client.post(
            "/analyze",
            data={
                "opinion_file": (BytesIO(b"%PDF-1.4 fake"), "brief.pdf"),
            },
            content_type="multipart/form-data",
        )

    assert response.status_code == 200
    assert b"Robert T. Hill" in response.data


def test_analyze_highlights_brief_language_passages(tmp_path: Path):
    app = create_app({"TESTING": True, "OPINION_DATA_DIR": tmp_path / "missing"})

    with app.test_client() as client:
        response = client.post(
            "/analyze",
            data={
                "opinion_file": (
                    BytesIO(
                        b"Opinion by Judge Elena M. Torres\n\nThe appellee's brief argued, \"The contract language was plain and unambiguous.\" The court adopted that wording."
                    ),
                    "opinion.txt",
                ),
            },
            content_type="multipart/form-data",
        )

    assert response.status_code == 200
    assert b"Brief language passages" in response.data
    assert b"The contract language was plain and unambiguous." in response.data
    assert b"Affinity Score" in response.data
    assert b"Writing Quality" in response.data


def test_analyze_uses_transcript_source_type(tmp_path: Path):
    opinions_dir = tmp_path / "opinions"
    transcripts_dir = tmp_path / "transcripts"
    opinions_dir.mkdir()
    transcripts_dir.mkdir()
    (transcripts_dir / "hearing_one.txt").write_text(
        "Opinion by Judge Elena M. Torres\n\nThe appellee's brief argued, \"The contract language was plain and unambiguous.\"",
        encoding="utf-8",
    )

    app = create_app(
        {
            "TESTING": True,
            "OPINION_DATA_DIR": opinions_dir,
            "TRANSCRIPT_DATA_DIR": transcripts_dir,
        }
    )

    with app.test_client() as client:
        response = client.post(
            "/analyze",
            data={"selected_transcript": "hearing_one"},
            content_type="multipart/form-data",
        )

    assert response.status_code == 200
    assert b"Transcript: Hearing One" in response.data
    assert b"transcript" in response.data
    assert b"Affinity Score" in response.data
    assert b"Writing Quality" in response.data


def test_analyze_handles_firm_document_upload_with_anonymization_and_audit_log(tmp_path: Path):
    app = create_app(
        {
            "TESTING": True,
            "OPINION_DATA_DIR": tmp_path / "opinions",
            "TRANSCRIPT_DATA_DIR": tmp_path / "transcripts",
            "FIRM_CONFIDENTIAL_DIR": tmp_path / "firm_confidential",
            "FIRM_CONFIDENTIAL_KEY_PATH": tmp_path / "firm_confidential.key",
            "AUDIT_LOG_PATH": tmp_path / "audit.log",
        }
    )

    with app.test_client() as client:
        response = client.post(
            "/analyze",
            data={
                "firm_file": (
                    BytesIO(
                        b"Confidential brief for Acme Corporation in Case No. 24-1234 seeking $2,500,000. Judge Jane Doe should review it."
                    ),
                    "firm_brief.txt",
                )
            },
            content_type="multipart/form-data",
        )

    assert response.status_code == 200
    assert b"Firm document: firm_brief.txt" in response.data
    assert b"Acme Corporation" not in response.data
    assert b"24-1234" not in response.data
    assert b"2,500,000" not in response.data
    assert b"[PROPER_NOUN]" in response.data or b"[CASE_NUMBER]" in response.data

    encrypted_files = list((tmp_path / "firm_confidential").glob("*.enc"))
    assert encrypted_files

    audit_log = (tmp_path / "audit.log").read_text(encoding="utf-8")
    assert "firm_upload" in audit_log
    assert "analysis_run" in audit_log


def test_load_folder_auto_assigns_known_judge_and_skips_duplicates(tmp_path: Path):
    opinions_dir = tmp_path / "opinions"
    transcripts_dir = tmp_path / "transcripts"
    folder_dir = tmp_path / "Maria Garcia Chambers"
    opinions_dir.mkdir()
    transcripts_dir.mkdir()
    folder_dir.mkdir()

    (folder_dir / "brief_one.txt").write_text("Opinion by Judge Maria Garcia\n\nThe court granted relief.", encoding="utf-8")
    (folder_dir / "brief_two.txt").write_text("Opinion by Judge Maria Garcia\n\nThe court granted relief.", encoding="utf-8")

    docx_path = folder_dir / "brief_three.docx"
    document = Document()
    document.add_paragraph("Opinion by Judge Maria Garcia")
    document.add_paragraph("The court granted relief.")
    document.save(docx_path)

    app = create_app(
        {
            "TESTING": True,
            "OPINION_DATA_DIR": opinions_dir,
            "TRANSCRIPT_DATA_DIR": transcripts_dir,
            "FIRM_CONFIDENTIAL_DIR": tmp_path / "firm_confidential",
            "FIRM_CONFIDENTIAL_KEY_PATH": tmp_path / "firm_confidential.key",
            "AUDIT_LOG_PATH": tmp_path / "audit.log",
            "PROCESSED_FILES_PATH": tmp_path / "processed_files.json",
        }
    )
    app.config["OPINION_SAMPLES"] = [type("Sample", (), {"judge_name": "Maria Garcia"})()]
    app.config["TRANSCRIPT_SAMPLES"] = []

    with app.test_client() as client:
        response = client.post(
            "/load-folder",
            data={
                "folder_path": str(folder_dir),
                "folder_context": "",
                "folder_judge_name": "",
            },
            content_type="multipart/form-data",
        )

    assert response.status_code == 302

    job_id = response.headers["Location"].rsplit("/", 1)[-1]
    deadline = time.time() + 5
    job = JOB_STORE.get(job_id)
    while job and job.status != "done" and time.time() < deadline:
        time.sleep(0.05)
        job = JOB_STORE.get(job_id)

    assert job is not None
    assert job.status == "done"
    assert job.judge_name == "Maria Garcia"
    assert job.processed_files >= 2
    assert job.skipped_files >= 1
    assert any(sample.source_type in {"opinion", "mixed"} for sample in job.created_samples)


def test_load_folder_marks_firm_sources_as_anonymized(tmp_path: Path):
    folder_dir = tmp_path / "Firm Winning Briefs"
    folder_dir.mkdir()
    (folder_dir / "brief_one.txt").write_text(
        "Acme Corporation paid $50,000 in Case No. 22-1144 before Judge Jane Doe.",
        encoding="utf-8",
    )

    app = create_app(
        {
            "TESTING": True,
            "OPINION_DATA_DIR": tmp_path / "opinions",
            "TRANSCRIPT_DATA_DIR": tmp_path / "transcripts",
            "FIRM_CONFIDENTIAL_DIR": tmp_path / "firm_confidential",
            "FIRM_CONFIDENTIAL_KEY_PATH": tmp_path / "firm_confidential.key",
            "AUDIT_LOG_PATH": tmp_path / "audit.log",
            "PROCESSED_FILES_PATH": tmp_path / "processed_files.json",
        }
    )

    with app.test_client() as client:
        response = client.post(
            "/load-folder",
            data={
                "folder_path": str(folder_dir),
                "folder_context": "",
                "folder_judge_name": "",
            },
            content_type="multipart/form-data",
        )

    assert response.status_code == 302

    job_id = response.headers["Location"].rsplit("/", 1)[-1]
    deadline = time.time() + 5
    job = JOB_STORE.get(job_id)
    while job and job.status != "done" and time.time() < deadline:
        time.sleep(0.05)
        job = JOB_STORE.get(job_id)

    assert job is not None
    assert job.status == "done"
    assert job.anonymized_texts
    assert "[PROPER_NOUN]" in job.anonymized_texts[0] or "[AMOUNT]" in job.anonymized_texts[0]
